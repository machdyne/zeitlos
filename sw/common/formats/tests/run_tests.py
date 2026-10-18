#!/usr/bin/env python3
"""
Zeitlos -- host tests for sw/common/formats (sw/apps/convert's library).

    python3 sw/common/formats/tests/run_tests.py [--no-lo]

Builds tests/zfhost.c with AddressSanitizer and UBSan, then:

  documents   fixtures made with python-docx, odfpy, ebooklib and by hand
              (rtf, html, md, txt in UTF-8, UTF-16 and Windows-1252) are
              converted to every document format; each result is read
              back by this library (as Markdown, compared against what is
              expected) and, for docx/odt, by python-docx/odfpy; every
              docx/odt/rtf/html written is opened by LibreOffice
              (headless) and its text compared
  tables      xlsx/ods/csv/tsv/zss fixtures to every table format, read
              back with openpyxl/odfpy and LibreOffice; formulas the
              sheet can evaluate survive, the rest arrive as values
  archives    zip, tar, tar.gz, gz extracted and compared byte for byte;
              archives with ../ and absolute paths refused, and nothing
              written outside the target
  hostile     every fixture truncated and corrupted at many points: the
              converter must fail cleanly (sanitizers on), never crash
  limits      a zip bomb stopped at ZF_ENTRY_MAX; cancelling stops at once

Requires: cc, python-docx, openpyxl, odfpy, ebooklib; LibreOffice for the
cross-checks (skipped with --no-lo or when soffice is missing).
"""

import os, sys, subprocess, shutil, tempfile, zipfile, tarfile, gzip, io, random, glob

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..', '..'))
COMMON = os.path.join(ROOT, 'sw', 'common')
TMP = tempfile.mkdtemp(prefix='zf_')
BIN = os.path.join(TMP, 'zfhost')
USE_LO = '--no-lo' not in sys.argv and shutil.which('soffice')

checks = 0
fails = 0

def ck(cond, what):
    global checks, fails
    checks += 1
    if not cond:
        fails += 1
        print('FAIL:', what)

def build():
    srcs = [os.path.join(HERE, 'zfhost.c')] + sorted(glob.glob(os.path.join(COMMON, 'formats', '*.c'))) + [
        os.path.join(COMMON, 'zinflate.c'), os.path.join(COMMON, 'zfix.c'),
        os.path.join(ROOT, 'sw', 'apps', 'sheet', 'sheet_core.c')]
    cmd = ['cc', '-std=gnu99', '-g', '-O1', '-Wall', '-Wextra', '-Werror', '-fsanitize=address,undefined',
           '-fno-sanitize-recover=undefined', '-I', COMMON, '-o', BIN] + srcs
    subprocess.check_call(cmd)

def run(*args, ok=True, env=None):
    p = subprocess.run([BIN] + list(args), capture_output=True, text=True, timeout=120,
                       env=dict(os.environ, ASAN_OPTIONS='detect_leaks=0'))
    if 'Sanitizer' in p.stderr or 'runtime error' in p.stderr:
        ck(False, 'sanitizer: %s\n%s' % (' '.join(args), p.stderr[-2000:]))
    if ok:
        ck(p.returncode == 0, 'zfhost %s -> %d %s' % (' '.join(args), p.returncode, p.stderr.strip()))
    return p

def f(name):
    return os.path.join(TMP, name)

def read(name):
    with open(f(name), encoding='utf-8') as h:
        return h.read()

def lo_text(path, fmt='txt:Text (encoded):UTF8'):
    out = os.path.join(TMP, 'lo_' + os.path.basename(path).replace('.', '_'))
    os.makedirs(out, exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', out, path],
                   capture_output=True, timeout=180)
    got = glob.glob(os.path.join(out, '*'))
    if not got:
        return None
    with open(got[0], encoding='utf-8-sig', errors='replace') as h:
        return h.read()

def lo_convert(path, ext):
    out = os.path.join(TMP, 'loc_' + ext)
    os.makedirs(out, exist_ok=True)
    subprocess.run(['soffice', '--headless', '--convert-to', ext, '--outdir', out, path],
                   capture_output=True, timeout=180)
    p = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + '.' + ext)
    return p if os.path.exists(p) else None

# -- documents --

EXPECT_MD = ('# The Title\n\n# Chapter One\n\n'
             'Plain text with **bold** and *italic* and «ünïcödé» – “quotes” 😀.\n\n'
             '- First bullet\n  - Second bullet\n\n## Sub heading\n\n'
             'Line one\x20\x20\nline two\tafter tab\n')

WORDS = ['The Title', 'Chapter One', 'Plain text with', 'bold', 'italic', '«ünïcödé»', '“quotes”',
         'First bullet', 'Second bullet', 'Sub heading', 'Line one', 'after tab', '<escape & me>']

def make_docs():
    import docx
    d = docx.Document()
    d.add_heading('The Title', 0)
    d.add_heading('Chapter One', 1)
    p = d.add_paragraph('Plain text with ')
    p.add_run('bold').bold = True
    p.add_run(' and ')
    p.add_run('italic').italic = True
    p.add_run(' and «ünïcödé» – “quotes” 😀.')
    d.add_paragraph('First bullet', style='List Bullet')
    d.add_paragraph('Second bullet', style='List Bullet 2')
    d.add_heading('Sub heading', 2)
    p = d.add_paragraph('Line one')
    p.add_run().add_break()
    p.add_run('line two\tafter tab')
    d.add_paragraph('<escape & me>')
    d.save(f('doc.docx'))

    with open(f('doc.html'), 'w', encoding='utf-8') as h:
        h.write('<!DOCTYPE html><html><head><title>T &amp; t</title><style>p{x:1}</style>'
                '<script>var a="<p>no</p>";</script></head><body>\n'
                '<h1>The   Title</h1><h1>Chapter\nOne</h1>'
                '<p>Plain text with <b>bold</b> and <em>italic</em> and &laquo;ünïcödé&raquo; &ndash; &#8220;quotes&#x201D; &#128512;.'
                '<ul><li>First bullet<ul><li><p>Second bullet</p></li></ul></li></ul>'
                '<h2>Sub heading</h2><p>Line one<br>line two after tab<!-- x --></p>'
                '<p>&lt;escape &amp; me&gt;</p></body></html>')

    with open(f('doc.md'), 'w', encoding='utf-8') as h:
        h.write('The Title\n=========\n\n# Chapter One #\n\nPlain text with **bold** and _italic_ and\n'
                '«ünïcödé» – “quotes” 😀.\n\n* First bullet\n  + Second bullet\n\n## Sub heading\n\n'
                'Line one  \nline two after tab\n\n```\ncode <x>\n  indented\n```\n\n'
                '\\<escape & me\\> and [a link](http://x) and `co*de*`\n')

    with open(f('doc.rtf'), 'w', encoding='ascii') as h:
        h.write('{\\rtf1\\ansi\\ansicpg1252\\deff0{\\fonttbl{\\f0\\fswiss Arial;}{\\f1\\fmodern Courier;}}'
                '{\\stylesheet{\\s0 Normal;}{\\s1\\b heading 1;}{\\s2 heading 2;}}'
                '{\\info{\\title Junk}}{\\*\\generator Hand;}\n'
                '\\pard\\s1 The Title\\par\n\\pard\\s1 Chapter One\\par\n'
                '\\pard Plain text with {\\b bold} and {\\i italic} and \\\'ab\\\'fcn\\\'efc\\\'f6d\\\'e9\\\'bb \\endash  \\ldblquote quotes\\rdblquote  \\uc1\\u-10179?\\u-8704?.\\par\n'
                '{\\pntext \\bullet\\tab}\\pard\\ls1\\ilvl0 First bullet\\par\n'
                '\\pard\\ls1\\ilvl1 Second bullet\\par\n'
                '\\pard\\outlinelevel1 Sub heading\\par\n'
                '\\pard Line one\\line line two\\tab after tab\\par\n'
                '{\\field{\\*\\fldinst HYPERLINK "x"}{\\fldrslt link text}}\\par\n'
                '\\pard\\f1 code line\\par\n'
                '\\pard\\f0 <escape & me> \\{braces\\}\\par}')

    with open(f('doc.txt'), 'w', encoding='utf-8', newline='') as h:
        h.write('First line\r\nSecond line ünï\r\n\r\nAfter blank\n')
    with open(f('doc16.txt'), 'wb') as h:
        h.write('\ufeffUTF-16 text ✓\r\nline 2\r\n'.encode('utf-16-le'))
    with open(f('doc1252.txt'), 'wb') as h:
        h.write('Caf\xe9 \x93quoted\x94 \x80 5\n'.encode('latin-1'))

    from odf.opendocument import OpenDocumentText
    from odf.style import Style, TextProperties
    from odf.text import H, P, Span, List, ListItem, S, Tab, LineBreak
    o = OpenDocumentText()
    b = Style(name='B', family='text'); b.addElement(TextProperties(fontweight='bold')); o.automaticstyles.addElement(b)
    i = Style(name='I', family='text'); i.addElement(TextProperties(fontstyle='italic')); o.automaticstyles.addElement(i)
    o.text.addElement(H(outlinelevel=1, text='The Title'))
    o.text.addElement(H(outlinelevel=1, text='Chapter One'))
    p = P(text='Plain text with ')
    p.addElement(Span(stylename=b, text='bold')); p.addText(' and ')
    p.addElement(Span(stylename=i, text='italic')); p.addText(' and «ünïcödé» – “quotes” 😀.')
    o.text.addElement(p)
    l = List(); li = ListItem(); li.addElement(P(text='First bullet'))
    l2 = List(); li2 = ListItem(); li2.addElement(P(text='Second bullet')); l2.addElement(li2); li.addElement(l2)
    l.addElement(li); o.text.addElement(l)
    o.text.addElement(H(outlinelevel=2, text='Sub heading'))
    p = P(text='Line one'); p.addElement(LineBreak()); p.addText('line two'); p.addElement(Tab()); p.addText('after tab')
    o.text.addElement(p)
    p = P(text='a'); p.addElement(S(c=3)); p.addText('b <escape & me>')
    o.text.addElement(p)
    o.save(f('doc.odt'))

    from ebooklib import epub
    b = epub.EpubBook()
    b.set_identifier('id1'); b.set_title('A Book'); b.set_language('en')
    c1 = epub.EpubHtml(title='One', file_name='text/c1.xhtml', content='<h1>Chapter One</h1><p>First <b>chapter</b>.</p>')
    c2 = epub.EpubHtml(title='Two', file_name='text/c2.xhtml', content='<h1>Chapter Two</h1><p>Second chapter.</p>')
    b.add_item(c1); b.add_item(c2)
    b.toc = (c1, c2)
    b.add_item(epub.EpubNcx()); b.add_item(epub.EpubNav())
    b.spine = [c2, c1]		# reading order differs from the manifest's
    epub.write_epub(f('book.epub'), b)

def test_docs():
    make_docs()
    # docx -> md is exactly what is expected
    run(f('doc.docx'), f('docx.md'))
    md = read('docx.md')
    ck(md.startswith(EXPECT_MD), 'docx -> md:\n' + md)

    # each source to md: the words are there, the structure is
    for src in ['doc.docx', 'doc.html', 'doc.md', 'doc.rtf', 'doc.odt']:
        out = src.replace('.', '_') + '.md'
        run(f(src), f(out))
        got = read(out)
        for w in WORDS[:12]:
            if src == 'doc.odt' and w == 'after tab':
                continue
            ck(w in got, '%s -> md missing %r:\n%s' % (src, w, got))
        ck('# Chapter One' in got, '%s: heading' % src)
        ck('**bold**' in got and '*italic*' in got, '%s: emphasis:\n%s' % (src, got))
        ck('- First bullet' in got and '  - Second bullet' in got, '%s: lists:\n%s' % (src, got))
        ck('## Sub heading' in got, '%s: heading 2' % src)
        ck('Line one  \nline two' in got, '%s: line break:\n%s' % (src, got))
    ck('code <x>\n  indented' in read('doc_md.md'), 'md fence kept')
    ck('<escape & me> and a link and co*de*' in read('doc_md.md'), 'md inline:\n' + read('doc_md.md'))
    ck('link text' in read('doc_rtf.md') and 'HYPERLINK' not in read('doc_rtf.md'), 'rtf field result only')
    ck('```\ncode line\n```' in read('doc_rtf.md'), 'rtf monospace as code:\n' + read('doc_rtf.md'))
    ck('{braces}' in read('doc_rtf.md') and 'Junk' not in read('doc_rtf.md'), 'rtf escapes, info skipped')
    ck('var a' not in read('doc_html.md') and 'x:1' not in read('doc_html.md'), 'html script/style skipped')
    ck('a   b' in read('doc_odt.md'), 'odt text:s kept:\n' + read('doc_odt.md'))

    # text encodings
    run(f('doc.txt'), f('txt.md'))
    ck(read('txt.md') == 'First line\n\nSecond line ünï\n\n\n\nAfter blank\n', 'txt lines:\n' + repr(read('txt.md')))
    run(f('doc16.txt'), f('t16.html'))
    ck('UTF-16 text ✓' in read('t16.html') and 'line 2' in read('t16.html'), 'utf-16')
    run(f('doc1252.txt'), f('t1252.txt'))
    ck(read('t1252.txt') == 'Café “quoted” € 5\n', 'cp1252: %r' % read('t1252.txt'))

    # epub: spine order, title
    run(f('book.epub'), f('book.html'))
    h = read('book.html')
    ck('<title>A Book</title>' in h, 'epub title')
    ck(h.find('Chapter Two') < h.find('Chapter One') and h.find('Chapter One') > 0, 'epub spine order')
    ck('<b>chapter</b>' in h, 'epub inline')

    # every source to every writer; the writers' output read back here
    # and by other software
    import docx
    from odf.opendocument import load
    from odf import teletype
    for src in ['doc.docx', 'doc.html', 'doc.md', 'doc.rtf', 'doc.odt', 'book.epub', 'doc.txt']:
        for ext in ['txt', 'md', 'html', 'rtf', 'docx', 'odt']:
            out = 'w_%s.%s' % (src.replace('.', '_'), ext)
            run(f(src), f(out))
            # read back
            run(f(out), f(out + '.md'))
            if src == 'doc.docx':
                back = read(out + '.md')
                if ext in ('md', 'html', 'docx', 'odt', 'rtf'):
                    want = EXPECT_MD.replace('\t', ' ') if ext == 'html' else EXPECT_MD
                    ck(back.startswith(want), 'round trip via %s:\n%s' % (ext, back))
            if ext == 'docx':
                try:
                    d = docx.Document(f(out))
                    ck(len(d.paragraphs) > 0, 'python-docx paragraphs ' + out)
                except Exception as e:
                    ck(False, 'python-docx %s: %s' % (out, e))
            if ext == 'odt':
                try:
                    o = load(f(out))
                    ck(len(teletype.extractText(o.text)) > 0, 'odfpy ' + out)
                except Exception as e:
                    ck(False, 'odfpy %s: %s' % (out, e))
            if USE_LO and src in ('doc.docx', 'book.epub') and ext in ('docx', 'odt', 'rtf', 'html'):
                t = lo_text(f(out))
                ck(t is not None, 'LibreOffice opens ' + out)
                if t and src == 'doc.docx':
                    for w in WORDS:
                        ck(w in t, 'LibreOffice %s missing %r' % (out, w))
                    ck('😀' in t, 'LibreOffice %s astral char' % out)

    # documents made by LibreOffice, read here
    if USE_LO:
        for ext in ['docx', 'odt', 'rtf', 'html']:
            p = lo_convert(f('doc.docx'), ext)
            ck(p is not None, 'LibreOffice made ' + ext)
            if not p:
                continue
            run(p, f('lo_in_%s.md' % ext))
            got = read('lo_in_%s.md' % ext)
            for w in WORDS[:12]:
                if ext == 'html' and w == 'after tab':
                    continue
                ck(w in got, 'from LibreOffice %s missing %r:\n%s' % (ext, w, got))
            ck('# Chapter One' in got, 'from LibreOffice %s: heading:\n%s' % (ext, got))
            ck('**bold**' in got, 'from LibreOffice %s: bold:\n%s' % (ext, got))
            ck('- First bullet' in got, 'from LibreOffice %s: list:\n%s' % (ext, got))

# -- tables --

def make_tabs():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Data'
    ws['A1'] = 'Name'; ws['B1'] = 'Value'
    ws['A2'] = 'x'; ws['B2'] = 1.5
    ws['A3'] = 'y, "q"'; ws['B3'] = 2
    ws['B4'] = '=SUM(B2:B3)'
    ws['C4'] = '=LEN(A2)'
    ws['A5'] = True
    ws['D6'] = 1.5e-7
    ws['A7'] = '12'
    ws['AB8'] = 'far'
    ws['B9'] = '=AVERAGE($B$2:B3)*2'
    ws['B10'] = '=IF(B2>1,B3^2,0)'
    wb.create_sheet('Second')['A1'] = 'not read'
    wb.save(f('t.xlsx'))
    with open(f('t.csv'), 'w', encoding='utf-8', newline='') as h:
        h.write('a,b,c\r\n1,"x, y","he said ""hi"""\r\n,,3.5\r\n"multi\nline",=A2,ü\r\n')
    with open(f('t.zss'), 'w') as h:
        h.write('zsheet 1\n!w A 12\nA1 Item\nB1 Cost\nA2 Pen\nB2 1.25\nA3 Ink\nB3 3\nB4 =SUM(B2:B3)\nB5 =B4/0\nA6 \'42\nC9 =AVG(B2:B3)\nC10 =$B$2*IF(B3>2,2,1)\n')

def test_tabs():
    make_tabs()
    import openpyxl
    from odf.opendocument import load
    from odf.table import TableRow, TableCell
    run(f('t.xlsx'), f('x.zss'))
    z = read('x.zss').splitlines()
    ck(z[0] == 'zsheet 1', 'zss header')
    ck('A1 Name' in z and 'B2 1.5' in z and 'A3 y, "q"' in z, 'xlsx -> zss cells:\n' + read('x.zss'))
    ck('B4 =SUM(B2:B3)' in z, 'formula kept')
    ck('C4 #N/A' in z or not any(l.startswith('C4 =') for l in z), 'unsupported formula as value')
    ck('A5 TRUE' in z, 'boolean')
    ck('D6 0.00000015' in z, 'exponent expanded: %s' % [l for l in z if l.startswith('D6')])
    ck("A7 '12" in z, 'numeric text forced')
    ck(not any(l.startswith('AB8') for l in z), 'beyond Z dropped')
    ck('B9 =AVERAGE($B$2:B3)*2' in z, 'absolute refs kept: %s' % [l for l in z if l.startswith('B9')])
    ck('B10 =IF(B2>1,B3^2,0)' in z, 'IF, comparison, power: %s' % [l for l in z if l.startswith('B10')])
    ck('not read' not in read('x.zss'), 'first sheet only')

    run(f('t.zss'), f('z.xlsx'))
    wb = openpyxl.load_workbook(f('z.xlsx'))
    ws = wb.active
    ck(ws['A1'].value == 'Item' and ws['B2'].value == 1.25 and ws['B4'].value == '=SUM(B2:B3)', 'zss -> xlsx')
    ck(ws['C9'].value == '=AVERAGE(B2:B3)', 'AVG written AVERAGE: %r' % ws['C9'].value)
    ck(ws['A6'].value == '42', 'forced text stays text')
    ck(ws['C10'].value == '=$B$2*IF(B3>2,2,1)', 'abs and IF to xlsx: %r' % ws['C10'].value)
    wbv = openpyxl.load_workbook(f('z.xlsx'), data_only=True)
    ck(wbv.active['B4'].value == 4.25, 'cached value: %r' % wbv.active['B4'].value)

    run(f('t.zss'), f('z.ods'))
    o = load(f('z.ods'))
    txt = str(o.spreadsheet)
    ck('Item' in txt and '4.25' in txt, 'zss -> ods')
    run(f('z.ods'), f('zo.zss'))
    back = read('zo.zss').splitlines()
    for line in ['A1 Item', 'B2 1.25', 'B4 =SUM(B2:B3)', "A6 '42", 'C9 =AVERAGE(B2:B3)', 'C10 =$B$2*IF(B3>2,2,1)']:
        ck(line in back, 'ods round trip %r:\n%s' % (line, read('zo.zss')))
    run(f('z.xlsx'), f('zx.zss'))
    back = read('zx.zss').splitlines()
    for line in ['A1 Item', 'B2 1.25', 'B4 =SUM(B2:B3)', "A6 '42"]:
        ck(line in back, 'xlsx round trip %r:\n%s' % (line, read('zx.zss')))

    run(f('t.csv'), f('c.xlsx'))
    ws = openpyxl.load_workbook(f('c.xlsx')).active
    ck(ws['B2'].value == 'x, y' and ws['C2'].value == 'he said "hi"' and ws['C3'].value == 3.5, 'csv quoting')
    ck(ws['A4'].value == 'multi\nline' and ws['C4'].value == 'ü', 'csv multiline, utf-8')
    run(f('c.xlsx'), f('c2.csv'))
    import csv
    rows = list(csv.reader(open(f('c2.csv'), encoding='utf-8', newline='')))
    ck(rows[1] == ['1', 'x, y', 'he said "hi"'] and rows[2] == ['', '', '3.5'], 'csv out: %r' % rows)
    run(f('t.csv'), f('c.tsv'))
    tsv = open(f('c.tsv'), encoding='utf-8', newline='').read()
    ck(tsv.split('\n')[1] == '1\tx, y\the said "hi"\r', 'tsv: %r' % tsv)

    if USE_LO:
        for out in ['z.xlsx', 'z.ods', 'c.xlsx']:
            t = lo_text(f(out), 'csv:Text - txt - csv (StarCalc):44,34,76')
            ck(t is not None and ('Item' in t or 'x, y' in t), 'LibreOffice opens %s: %r' % (out, t))
            if out == 'z.xlsx' and t:
                ck('4.25' in t, 'LibreOffice evaluates formula')
        for ext in ['xlsx', 'ods']:
            p = lo_convert(f('t.xlsx'), ext)
            run(p, f('lo_%s.zss' % ext))
            z = read('lo_%s.zss' % ext).splitlines()
            ck('B4 =SUM(B2:B3)' in z and 'B2 1.5' in z, 'from LibreOffice %s:\n%s' % (ext, '\n'.join(z)))

# -- archives --

def test_archives():
    src = f('arch_src')
    os.makedirs(os.path.join(src, 'sub', 'deep'), exist_ok=True)
    files = {'a.txt': b'hello\n', 'sub/b.bin': bytes(range(256)) * 300, 'sub/deep/c.txt': b'x' * 70000, 'empty': b''}
    for k, v in files.items():
        with open(os.path.join(src, k), 'wb') as h:
            h.write(v)
    with zipfile.ZipFile(f('a.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
        for k in files:
            z.write(os.path.join(src, k), k)
        z.writestr('dir/', '')
    with tarfile.open(f('a.tar'), 'w') as t:
        for k in files:
            t.add(os.path.join(src, k), k)
        long = 'l' * 120 + '/' + 'n' * 120 + '.txt'
        ti = tarfile.TarInfo(long); ti.size = 3
        t.addfile(ti, io.BytesIO(b'abc'))
    with tarfile.open(f('a.tar.gz'), 'w:gz') as t:
        for k in files:
            t.add(os.path.join(src, k), k)
    with gzip.open(f('one.txt.gz'), 'wb') as h:
        h.write(files['sub/b.bin'])

    for a in ['a.zip', 'a.tar', 'a.tar.gz']:
        out = f('x_' + a.replace('.', '_'))
        run(f(a), out)
        for k, v in files.items():
            p = os.path.join(out, k)
            ck(os.path.exists(p) and open(p, 'rb').read() == v, '%s: %s' % (a, k))
    ck(os.path.exists(os.path.join(f('x_a_tar'), 'l' * 120, 'n' * 120 + '.txt')), 'tar long name')
    ck(os.path.isdir(os.path.join(f('x_a_zip'), 'dir')), 'zip directory entry')
    run(f('one.txt.gz'), f('one.txt'))
    ck(open(f('one.txt'), 'rb').read() == files['sub/b.bin'], 'gz')

    # hostile paths
    for name in ['../evil.txt', '/abs.txt', 'a/../../evil.txt', 'C:\\x.txt', '..\\evil.txt']:
        with zipfile.ZipFile(f('bad.zip'), 'w') as z:
            z.writestr('ok.txt', 'fine')
            z.writestr(name, 'evil')
        out = f('x_bad')
        shutil.rmtree(out, ignore_errors=True)
        p = run(f('bad.zip'), out, ok=False)
        ck(p.returncode != 0, 'zip path %r refused' % name)
        ck(not os.path.exists(os.path.join(TMP, 'evil.txt')) and not os.path.exists('/abs.txt'), 'nothing escaped')
        ck(not os.path.exists(os.path.join(out, 'ok.txt')), 'nothing extracted from a bad zip')
        with tarfile.open(f('bad.tar'), 'w') as t:
            ti = tarfile.TarInfo(name.replace('\\', '/')); ti.size = 4
            t.addfile(ti, io.BytesIO(b'evil'))
        p = run(f('bad.tar'), f('x_badtar'), ok=False)
        if name.startswith('C:'):
            continue
        ck(p.returncode != 0, 'tar path %r refused' % name)
    ck(not os.path.exists(os.path.join(TMP, 'evil.txt')), 'nothing escaped (tar)')

    # a bomb: one entry claiming small, expanding past ZF_ENTRY_MAX
    with zipfile.ZipFile(f('bomb.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('big.bin', b'\0' * (70 * 1024 * 1024))
    p = run(f('bomb.zip'), f('x_bomb'), ok=False)
    ck(p.returncode != 0 and 'too large' in p.stderr, 'bomb stopped: ' + p.stderr)

# -- hostile input --

def test_hostile():
    rnd = random.Random(7)
    srcs = ['doc.docx', 'doc.odt', 'doc.rtf', 'doc.html', 'doc.md', 'book.epub', 't.xlsx', 'z.ods', 't.csv',
            't.zss', 'a.zip', 'a.tar', 'a.tar.gz', 'one.txt.gz']
    n = 0
    for s in srcs:
        data = open(f(s), 'rb').read()
        ext = s.split('.', 1)[1]
        outext = {'docx': 'md', 'odt': 'docx', 'rtf': 'odt', 'html': 'rtf', 'md': 'html', 'epub': 'txt',
                  'xlsx': 'ods', 'ods': 'xlsx', 'csv': 'zss', 'zss': 'csv'}.get(ext.split('.')[-1], None)
        cases = []
        for cut in sorted(set([0, 1, 4, 30, 100, len(data) // 2, len(data) - 22, len(data) - 1])):
            if 0 <= cut < len(data):
                cases.append(data[:cut])
        for _ in range(40):
            b = bytearray(data)
            for _ in range(rnd.randint(1, 8)):
                b[rnd.randrange(len(b))] = rnd.randrange(256)
            cases.append(bytes(b))
        for i, c in enumerate(cases):
            p = f('h_' + str(i) + '_' + s)
            with open(p, 'wb') as h:
                h.write(c)
            if outext:
                run('-t', outext, p, f('h_out.' + outext), ok=False)
            else:
                shutil.rmtree(f('h_x'), ignore_errors=True)
                run(p, f('h_x'), ok=False)
            os.unlink(p)
            n += 1
    ck(n > 500, 'hostile cases run: %d' % n)

def test_cancel():
    # cancelling: stops with "cancelled" at the first poll after the request
    for src, out in [('doc.docx', 'c.odt'), ('t.xlsx', 'c.ods'), ('a.tar.gz', 'c_x'), ('book.epub', 'c.md')]:
        p = run('-c', '3', f(src), f(out), ok=False)
        ck(p.returncode == 1 and 'cancelled' in p.stderr, 'cancel %s: %s' % (src, p.stderr))
    # a cancelled extraction: no half-written file
    shutil.rmtree(f('cx'), ignore_errors=True)
    p = run('-c', '6', f('a.tar'), f('cx'), ok=False)
    ck('cancelled' in p.stderr, 'tar cancel: ' + p.stderr)
    big = os.path.join(f('cx'), 'sub', 'deep', 'c.txt')
    ck(not os.path.exists(big) or os.path.getsize(big) == 70000, 'no partial file left')
    p = run('-m', '2000', f('book.epub'), f('m.txt'), ok=False)
    ck('out of memory' in p.stderr, 'heap too small reported: ' + p.stderr)

def main():
    build()
    test_docs()
    test_tabs()
    test_archives()
    test_hostile()
    test_cancel()
    print('%d checks, %d failed%s' % (checks, fails, '' if USE_LO else ' (LibreOffice checks skipped)'))
    shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(1 if fails else 0)

if __name__ == '__main__':
    main()
